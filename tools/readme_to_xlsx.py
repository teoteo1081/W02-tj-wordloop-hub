#!/usr/bin/env python3
"""Sinh file Excel XEM công việc từ README.md mục "Việc còn dang dở" + git log (TJ 2026-10-06: "mình muốn xem excel").
KHÔNG phải nguồn sự thật — nguồn thật là README.md (xem TEAM_PROCESS.md mục 3). Chạy lại bất cứ lúc nào để làm mới:
    python3 tools/readme_to_xlsx.py [đường_dẫn_ra.xlsx]      (cần: pip install openpyxl)
Sau khi chạy nên chạy recalc của skill xlsx (nếu có) để các ô công thức ở sheet "Tóm tắt" có sẵn số."""
import re, subprocess, sys, datetime, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "TJ_CongViec.xlsx")
md = (ROOT / "README.md").read_text(encoding="utf8")
sec = md[md.index("## Việc còn dang dở"):]
nxt = re.search(r"\n## ", sec[5:]); sec = sec[: nxt.start() + 5] if nxt else sec

# --- tách các gạch đầu dòng cấp 1 ---
items, cur = [], None
for line in sec.splitlines()[1:]:
    if line.startswith("- **"):
        cur = [line[2:]]; items.append(cur)
    elif cur is not None and (line.startswith(" ") or line.strip() == ""):
        cur.append(line.strip())
    elif line.startswith(">") or line.startswith("#"):
        cur = None

def classify(i, title, body):
    t, b = title, body
    if i == 0: return "Cao", "Đang làm / ưu tiên"
    if "TOEIC LISTENING" in t: return "Thấp", "Chờ TJ duyệt mẫu"
    if "NHẮC TJ" in t: return "Trung bình", "Chờ TJ"
    if t.startswith("✅"): return "—", "Đã xong"
    if re.search(r"(Bảng|game|Game|gamelayer)\s*v\d{2,3}|\bv\d{3}\b", t): return "—", "Nhật ký bản game (đã làm)"
    if "CHƯA CODE" in t or "CHƯA CODE" in b[:400] or "đang bàn" in t or "Ý TƯỞNG" in t or "ý tưởng" in t[:60]: return "Chưa xếp", "Ý tưởng / đang bàn, chưa làm"
    if "ĐANG LÀM" in t or "▶" in t: return "Trung bình", "Đang làm"
    return "Chưa xếp", "Ghi chú / theo dõi"

rows = []
for i, it in enumerate(items):
    first = it[0]
    m = re.match(r"\*\*(.+?)\*\*:?\s*(.*)", first, re.S)
    title, rest = (m.group(1), m.group(2)) if m else (first[:80], first)
    body = (rest + "\n" + "\n".join(x for x in it[1:] if x)).strip()
    d = re.search(r"(\d{4}-\d{2}-\d{2})", title) or re.search(r"(\d{4}-\d{2}-\d{2})", body[:200])
    pr, st = classify(i, title, body)
    flat = re.sub(r"\s+", " ", body)
    rows.append([len(rows) + 1, pr, st, re.sub(r"[*`]", "", title)[:150], d.group(1) if d else "", re.sub(r"[*`]", "", flat)[:300] + ("…" if len(flat) > 300 else ""), re.sub(r"[*`]", "", flat)[:32000]])

log = subprocess.run(["git", "-C", str(ROOT), "log", "-40", "--date=short", "--pretty=%ad|%h|%s"], capture_output=True, text=True).stdout.splitlines()
done = [l.split("|", 2) for l in log if l.count("|") >= 2]

F = "Arial"; hdr_fill = PatternFill("solid", fgColor="1F3864"); thin = Side(style="thin", color="BBBBBB")
def head(ws, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(1, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = hdr_fill
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 30

wb = Workbook()
ws = wb.active; ws.title = "Việc còn dang dở"
head(ws, ["#", "Ưu tiên", "Trạng thái", "Tiêu đề", "Ngày", "Tóm tắt (300 ký tự đầu)", "Chi tiết đầy đủ (bấm vào ô để đọc)"], [5, 12, 20, 45, 12, 70, 60])
fills = {"Cao": "F8CBAD", "Trung bình": "FFE699", "Thấp": "C6E0B4", "Chưa xếp": "EDEDED", "—": "DDEBF7"}
for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        x = ws.cell(r, c, v); x.font = Font(name=F, size=10)
        x.alignment = Alignment(vertical="top", wrap_text=(c in (4, 6)))
        x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.cell(r, 2).fill = PatternFill("solid", fgColor=fills.get(row[1], "FFFFFF"))
    ws.row_dimensions[r].height = 75
ws.auto_filter.ref = f"A1:G{len(rows)+1}"

w2 = wb.create_sheet("Đã xong gần đây")
head(w2, ["Ngày", "Mã commit", "Nội dung (git log, 40 commit mới nhất)"], [12, 12, 120])
for r, (d, h, s) in enumerate(done, 2):
    for c, v in enumerate([d, h, s], 1):
        x = w2.cell(r, c, v); x.font = Font(name=F, size=10); x.alignment = Alignment(vertical="top", wrap_text=(c == 3))
w2.auto_filter.ref = f"A1:C{len(done)+1}"

w3 = wb.create_sheet("Tóm tắt", 0)
w3["A1"] = "TJ WordLoop Hub — bảng xem công việc"; w3["A1"].font = Font(name=F, bold=True, size=14)
w3["A2"] = f"Sinh tự động ngày {datetime.date.today().isoformat()} từ README.md (mục 'Việc còn dang dở') và git log."; w3["A2"].font = Font(name=F, italic=True, size=10)
n = len(rows) + 1
tbl = [("Ưu tiên", "Số mục"), ("Cao", f'=COUNTIF(\'Việc còn dang dở\'!$B$2:$B${n},A5)'), ("Trung bình", f'=COUNTIF(\'Việc còn dang dở\'!$B$2:$B${n},A6)'),
       ("Thấp", f'=COUNTIF(\'Việc còn dang dở\'!$B$2:$B${n},A7)'), ("Chưa xếp", f'=COUNTIF(\'Việc còn dang dở\'!$B$2:$B${n},A8)'), ("—", f'=COUNTIF(\'Việc còn dang dở\'!$B$2:$B${n},A9)'), ("Tổng", "=SUM(B5:B9)")]
for k, (a, b) in enumerate(tbl, 4):
    for c, v in enumerate((a, b), 1):
        x = w3.cell(k, c, v); x.font = Font(name=F, bold=(k in (4, 10)), size=11)
        if k == 4: x.fill = hdr_fill; x.font = Font(name=F, bold=True, color="FFFFFF")
w3["A12"] = "Cách đọc"; w3["A12"].font = Font(name=F, bold=True, size=12)
notes = ["• File này CHỈ ĐỂ XEM. Nguồn thật là README.md trong GitHub (https://github.com/teoteo1081/W02-tj-wordloop-hub). Sửa việc ở README, rồi chạy lại: python3 tools/readme_to_xlsx.py",
         "• 'Ưu tiên' và 'Trạng thái' do script đoán từ tiêu đề/biểu tượng trong README (mục đầu = Cao; mục TOEIC Listening = Thấp; ✅ = Đã xong). Có thể sai — nội dung trong README mới đúng.",
         "• Mục 'Chưa xếp' = chưa ai gán độ ưu tiên; hỏi TJ trước khi làm.",
         "• Sheet 'Đã xong gần đây' = 40 commit mới nhất của repo (git log)."]
for k, t in enumerate(notes, 13):
    w3.cell(k, 1, t).font = Font(name=F, size=10)
w3.column_dimensions["A"].width = 24; w3.column_dimensions["B"].width = 12
wb.calculation.fullCalcOnLoad = True
wb.save(OUT); print("đã lưu", OUT, "·", len(rows), "mục việc ·", len(done), "commit")
